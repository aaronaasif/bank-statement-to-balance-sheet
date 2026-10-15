"""
Balance Sheet Builder
---------------------
Upload a bank statement PDF (S.No / Date / Particulars / Inst. No. / Memo /
Debit / Credit / Balance layout, newest transaction first) and get a balance
sheet with one tab per category.

Run:   uvicorn balance_app:app --reload
Open:  http://127.0.0.1:8000

Privacy: the uploaded PDF is read straight from the request into memory and is
never written to disk. It is discarded as soon as it has been parsed. Only the
generated Excel file is kept in memory for RESULT_TTL_SECONDS so it can be
downloaded, then it is thrown away too.
"""
import io
import re
import secrets
import time
from decimal import Decimal, ROUND_HALF_UP

import pdfplumber
from fastapi import FastAPI, HTTPException, Request
from fastapi.responses import HTMLResponse, JSONResponse, Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from starlette.concurrency import run_in_threadpool

# ----------------------------------------------------------------------------
# Settings you may want to change
# ----------------------------------------------------------------------------
MAX_UPLOAD_MB = 20
RESULT_TTL_SECONDS = 15 * 60          # how long the Excel download stays available

WHT_RATE = Decimal("0.20")            # withholding tax on gross profit
PROFIT_KEYWORDS = ["PROFIT"]          # only used for the oldest row, which has no row below it

# First matching keyword wins, so keep the order. Gross profit rows are detected
# separately (from the 20% WHT gap in the running balance).
RULES = [
    ("UTILITY BILL", "Utility bills"),
    ("CASH WITHDRAWAL", "Cash withdrawals"),
    ("WITHDRL", "Other withdrawals"),
    ("POS PURCHASE", "POS"),
    ("RAAST", "Raast"),
    ("IBT", "Transfers"),
    ("DEPOSIT", "Deposits"),
]

ZERO = Decimal("0.00")
CENT = Decimal("0.01")
AMOUNT = re.compile(r"^[\d,]+\.\d{2}$")
DATE = re.compile(r"^\d{4}-\d{2}-\d{2}$")


class StatementError(Exception):
    """A problem with the uploaded statement that we can explain to the user."""


# ----------------------------------------------------------------------------
# 1. Reading the PDF
# ----------------------------------------------------------------------------
def num(text):
    return Decimal(text.replace(",", ""))


def wht_of(amount):
    return (amount * WHT_RATE).quantize(CENT, rounding=ROUND_HALF_UP)


def group_lines(words, tol=3):
    """Group words that sit on the same visual line."""
    lines = []
    for w in sorted(words, key=lambda w: (w["top"], w["x0"])):
        if lines and abs(w["top"] - lines[-1][0]) <= tol:
            lines[-1][1].append(w)
        else:
            lines.append([w["top"], [w]])
    return [sorted(ws, key=lambda w: w["x0"]) for _, ws in lines]


def find_bounds(lines):
    """Find the header line and return the x positions that separate the columns."""
    for ws in lines:
        h = {w["text"]: w for w in ws}
        if all(k in h for k in ("Debit", "Credit", "Balance")):
            left = h["Memo"]["x1"] if "Memo" in h else h["Debit"]["x0"] - 40
            b1 = (left + h["Debit"]["x0"]) / 2
            b2 = (h["Debit"]["x1"] + h["Credit"]["x0"]) / 2
            b3 = (h["Credit"]["x1"] + h["Balance"]["x0"]) / 2
            return b1, b2, b3
    return None


def read_pdf(data):
    """Return (rows, page_count). Rows are in statement order (newest first)."""
    try:
        pdf = pdfplumber.open(io.BytesIO(data))
    except Exception:
        raise StatementError("This file could not be opened as a PDF. It may be damaged or password protected.")

    rows = []
    bounds = None
    last_page = last_top = None

    with pdf:
        pages = len(pdf.pages)
        for p, page in enumerate(pdf.pages):
            lines = group_lines(page.extract_words())
            bounds = find_bounds(lines) or bounds
            if not bounds:
                continue
            b1, b2, b3 = bounds

            for ws in lines:
                top = ws[0]["top"]
                if len(ws) > 1 and ws[0]["text"].isdigit() and DATE.match(ws[1]["text"]):
                    debit, credit, balance = ZERO, ZERO, ZERO
                    desc = []
                    for w in ws[2:]:
                        cx = (w["x0"] + w["x1"]) / 2
                        if AMOUNT.match(w["text"]) and cx >= b1:
                            v = num(w["text"])
                            if cx < b2:
                                debit = v
                            elif cx < b3:
                                credit = v
                            else:
                                balance = v
                        elif cx < b1:
                            desc.append(w["text"])
                    rows.append({"sno": int(ws[0]["text"]), "date": ws[1]["text"],
                                 "desc": " ".join(desc), "debit": debit,
                                 "credit": credit, "balance": balance, "wht": ZERO})
                    last_page, last_top = p, top
                elif (rows and last_page == p and top - last_top < 14
                      and all((w["x0"] + w["x1"]) / 2 < b1 for w in ws)):
                    # description that wrapped onto a second line
                    rows[-1]["desc"] += " " + " ".join(w["text"] for w in ws)
                    last_top = top

    if not rows:
        raise StatementError(
            "No transactions were found. This app reads text-based PDFs in the "
            "S.No / Date / Particulars / Debit / Credit / Balance layout. "
            "A scanned (image) statement or a different layout will not work."
        )
    # Make sure the newest transaction comes first.
    if rows[0]["date"] < rows[-1]["date"]:
        rows.reverse()
    return rows, pages


# ----------------------------------------------------------------------------
# 2. Working out WHT, categories and totals
# ----------------------------------------------------------------------------
def apply_wht(rows):
    """
    The bank credits only 80% of each gross profit payment and does not show the
    20% withholding tax as a line. Rows are newest first, so the older
    transaction is the row below. If a credit raised the balance by less than
    the credit, and the gap is exactly 20% of the credit, that row is gross profit.
    """
    for i, r in enumerate(rows):
        r["wht"] = ZERO
        if r["credit"] <= 0 or r["debit"] > 0:
            continue
        if i + 1 < len(rows):
            change = r["balance"] - rows[i + 1]["balance"]
            gap = r["credit"] - change
            if gap > 0 and abs(gap - wht_of(r["credit"])) <= CENT:
                r["wht"] = gap
        elif any(k in r["desc"].upper() for k in PROFIT_KEYWORDS):
            # oldest row: nothing below it to compare with, so go by the description
            r["wht"] = wht_of(r["credit"])


def categorize(r):
    if r["wht"] > 0:
        return "Gross profit"
    d = r["desc"].upper()
    for key, cat in RULES:
        if key in d:
            return cat
    return "Other"


def net_of(r):
    return r["credit"] - r["debit"] - r["wht"]


def analyse(rows, pages):
    apply_wht(rows)
    for r in rows:
        r["category"] = categorize(r)

    # every row must fit the running balance
    unexplained = []
    for i in range(len(rows) - 1):
        r = rows[i]
        change = r["balance"] - rows[i + 1]["balance"]
        if change != net_of(r):
            unexplained.append({"sno": r["sno"], "date": r["date"], "desc": r["desc"],
                                "difference": float(change - net_of(r))})

    opening = rows[-1]["balance"] - net_of(rows[-1])
    closing = rows[0]["balance"]
    calculated = opening + sum(net_of(r) for r in rows)
    total_credit = sum(r["credit"] for r in rows)
    total_debit = sum(r["debit"] for r in rows)
    total_wht = sum(r["wht"] for r in rows)
    gross_profit = sum(r["credit"] for r in rows if r["wht"] > 0)

    # monthly summary, oldest month first
    months = {}
    for r in reversed(rows):
        m = months.setdefault(r["date"][:7], {
            "month": r["date"][:7], "gross_profit": ZERO, "credits": ZERO,
            "debits": ZERO, "wht": ZERO, "net": ZERO, "closing_balance": ZERO})
        m["credits"] += r["credit"]
        m["debits"] += r["debit"]
        m["wht"] += r["wht"]
        m["net"] += net_of(r)
        if r["wht"] > 0:
            m["gross_profit"] += r["credit"]
        m["closing_balance"] = r["balance"]       # last one seen is the latest in the month
    monthly = list(months.values())

    # category summary, in a fixed, readable order
    order = ["Gross profit"] + [c for _, c in RULES] + ["Other"]
    cats = {}
    for r in rows:
        c = cats.setdefault(r["category"], {"name": r["category"], "count": 0,
                                            "debits": ZERO, "credits": ZERO, "wht": ZERO})
        c["count"] += 1
        c["debits"] += r["debit"]
        c["credits"] += r["credit"]
        c["wht"] += r["wht"]
    categories = [cats[n] for n in order if n in cats]

    return {
        "rows": rows, "monthly": monthly, "categories": categories, "pages": pages,
        "opening": opening, "closing": closing, "calculated": calculated,
        "credits": total_credit, "debits": total_debit, "wht": total_wht,
        "gross_profit": gross_profit, "unexplained": unexplained,
        "reconciled": (calculated == closing and not unexplained),
        "period_from": min(r["date"] for r in rows),
        "period_to": max(r["date"] for r in rows),
    }


# ----------------------------------------------------------------------------
# 3. Excel file
# ----------------------------------------------------------------------------
def build_workbook(a):
    wb = Workbook()
    wb.remove(wb.active)
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="0F6B5C")
    used = set()

    def sheet_name(name):
        base = re.sub(r"[\[\]:*?/\\]", "", name)[:31] or "Sheet"
        n, k = base, 2
        while n in used:
            n = f"{base[:28]} {k}"
            k += 1
        used.add(n)
        return n

    def add(name, headers, data, money_cols, total=None):
        ws = wb.create_sheet(sheet_name(name))
        ws.append(headers)
        for c in ws[1]:
            c.font, c.fill = head_font, head_fill
            c.alignment = Alignment(horizontal="center")
        for row in data:
            ws.append([float(v) if isinstance(v, Decimal) else v for v in row])
        if total:
            ws.append([float(v) if isinstance(v, Decimal) else v for v in total])
            for c in ws[ws.max_row]:
                c.font = Font(bold=True)
        for idx in money_cols:
            for c in ws[get_column_letter(idx)][1:]:
                c.number_format = "#,##0.00"
        for col in ws.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(width + 3, 45)
        ws.freeze_panes = "A2"

    add("Monthly",
        ["Month", "Gross profit", "Credits", "Debits", "WHT", "Net movement", "Closing balance"],
        [[m["month"], m["gross_profit"], m["credits"], m["debits"], m["wht"], m["net"], m["closing_balance"]]
         for m in a["monthly"]],
        [2, 3, 4, 5, 6, 7],
        ["Total", a["gross_profit"], a["credits"], a["debits"], a["wht"],
         sum(m["net"] for m in a["monthly"]), None])

    add("Summary",
        ["Category", "Transactions", "Debits", "Credits", "WHT"],
        [[c["name"], c["count"], c["debits"], c["credits"], c["wht"]] for c in a["categories"]],
        [3, 4, 5],
        ["Total", len(a["rows"]), a["debits"], a["credits"], a["wht"]])

    tx_headers = ["Date", "Description", "Debit", "Credit", "WHT", "Balance"]
    for c in a["categories"]:
        rows = [r for r in a["rows"] if r["category"] == c["name"]]
        add(c["name"], tx_headers,
            [[r["date"], r["desc"], r["debit"], r["credit"], r["wht"], r["balance"]] for r in rows],
            [3, 4, 5, 6],
            ["", f"Total ({len(rows)} transactions)", c["debits"], c["credits"], c["wht"], None])

    add("All transactions", ["Date", "Description", "Category", "Debit", "Credit", "WHT", "Balance"],
        [[r["date"], r["desc"], r["category"], r["debit"], r["credit"], r["wht"], r["balance"]]
         for r in a["rows"]],
        [4, 5, 6, 7])

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ----------------------------------------------------------------------------
# 4. Web app
# ----------------------------------------------------------------------------
app = FastAPI(title="Balance Sheet Builder")

_downloads = {}   # token -> (expires_at, xlsx bytes). Memory only.


def _purge():
    now = time.time()
    for t in [t for t, (exp, _) in _downloads.items() if exp < now]:
        del _downloads[t]
    while len(_downloads) > 50:                   # safety cap
        del _downloads[next(iter(_downloads))]


def _f(d):
    return float(d)


def _json_result(a, token):
    return {
        "token": token,
        "summary": {
            "pages": a["pages"], "count": len(a["rows"]),
            "period_from": a["period_from"], "period_to": a["period_to"],
            "opening": _f(a["opening"]), "closing": _f(a["closing"]),
            "calculated": _f(a["calculated"]), "credits": _f(a["credits"]),
            "debits": _f(a["debits"]), "wht": _f(a["wht"]),
            "gross_profit": _f(a["gross_profit"]),
            "reconciled": a["reconciled"], "unexplained": a["unexplained"],
        },
        "monthly": [{**m, **{k: _f(m[k]) for k in
                             ("gross_profit", "credits", "debits", "wht", "net", "closing_balance")}}
                    for m in a["monthly"]],
        "categories": [{**c, **{k: _f(c[k]) for k in ("debits", "credits", "wht")}}
                       for c in a["categories"]],
        "transactions": [{"sno": r["sno"], "date": r["date"], "desc": r["desc"],
                          "category": r["category"], "debit": _f(r["debit"]),
                          "credit": _f(r["credit"]), "wht": _f(r["wht"]),
                          "balance": _f(r["balance"])} for r in a["rows"]],
    }


def process_pdf(data):
    rows, pages = read_pdf(data)
    a = analyse(rows, pages)
    return a, build_workbook(a)


@app.post("/api/process")
async def process(request: Request):
    """Receives the raw PDF as the request body. Nothing is written to disk."""
    limit = MAX_UPLOAD_MB * 1024 * 1024
    declared = request.headers.get("content-length")
    if declared and declared.isdigit() and int(declared) > limit:
        raise HTTPException(413, f"File is too large (limit {MAX_UPLOAD_MB} MB).")

    data = bytearray()
    async for chunk in request.stream():
        data.extend(chunk)
        if len(data) > limit:
            raise HTTPException(413, f"File is too large (limit {MAX_UPLOAD_MB} MB).")
    data = bytes(data)

    if not data.startswith(b"%PDF-"):
        raise HTTPException(400, "Please upload a PDF file.")

    try:
        a, xlsx = await run_in_threadpool(process_pdf, data)
    except StatementError as e:
        raise HTTPException(422, str(e))
    finally:
        del data                                   # the statement is no longer needed

    _purge()
    token = secrets.token_urlsafe(24)
    _downloads[token] = (time.time() + RESULT_TTL_SECONDS, xlsx)
    return JSONResponse(_json_result(a, token), headers={"Cache-Control": "no-store"})


@app.get("/api/download/{token}")
def download(token: str):
    _purge()
    item = _downloads.get(token)
    if not item:
        raise HTTPException(404, "This download has expired. Please upload the statement again.")
    return Response(
        item[1],
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="balance_sheet.xlsx"',
                 "Cache-Control": "no-store"},
    )


@app.get("/", response_class=HTMLResponse)
def home():
    return PAGE


PAGE = r"""<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Balance Sheet Builder</title>
<style>
  :root {
    --bg: #f4f1ec; --card: #ffffff; --ink: #1d2433; --muted: #6b7385; --line: #e4e0d8;
    --accent: #0f6b5c; --accent-soft: #e2f1ee; --debit: #b4412f; --credit: #0f6b5c;
    --warn: #8a5a00; --warn-soft: #fdf1d6;
  }
  * { box-sizing: border-box; margin: 0; padding: 0; }
  body { background: var(--bg); color: var(--ink); font-family: Inter, system-ui, -apple-system, "Segoe UI", sans-serif; font-size: 15px; }
  .topbar { display: flex; align-items: center; padding: 20px 32px; background: var(--card); border-bottom: 1px solid var(--line); }
  .logo { display: flex; align-items: center; gap: 12px; font-weight: 700; font-size: 20px; letter-spacing: -0.2px; }
  .logo i { width: 34px; height: 34px; border-radius: 9px; background: var(--accent); display: inline-block; position: relative; }
  .logo i:after { content: ""; position: absolute; left: 9px; top: 9px; width: 16px; height: 16px; border-left: 3px solid #fff; border-bottom: 3px solid #fff; }
  .wrap { max-width: 1312px; margin: 0 auto; padding: 0 32px 56px; }
  h1 { font-size: 32px; letter-spacing: -0.8px; margin: 36px 0 6px; }
  .sub { color: var(--muted); font-size: 16px; margin-bottom: 26px; }

  .upload { display: grid; grid-template-columns: 1fr 280px; gap: 20px; }
  .drop { background: var(--card); border: 2px dashed #bfc9c5; border-radius: 16px; padding: 24px 28px; display: flex; align-items: center; gap: 22px; cursor: pointer; flex-wrap: wrap; }
  .drop.over { border-color: var(--accent); background: var(--accent-soft); }
  .drop .ico { width: 60px; height: 60px; border-radius: 14px; background: var(--accent-soft); position: relative; flex: none; }
  .drop .ico:before { content: ""; position: absolute; left: 28px; top: 17px; width: 3px; height: 26px; background: var(--accent); }
  .drop .ico:after { content: ""; position: absolute; left: 22px; top: 17px; width: 12px; height: 12px; border-left: 3px solid var(--accent); border-top: 3px solid var(--accent); transform: translate(3px, 2px) rotate(45deg); }
  .drop h3 { font-size: 18px; margin-bottom: 4px; }
  .drop p { color: var(--muted); }
  .file { margin-left: auto; display: flex; align-items: center; gap: 12px; background: var(--bg); border: 1px solid var(--line); border-radius: 12px; padding: 10px 14px; }
  .file[hidden] { display: none; }
  .file .pdf { background: var(--debit); color: #fff; font-size: 11px; font-weight: 700; padding: 6px 7px; border-radius: 6px; }
  .file small { color: var(--muted); display: block; font-size: 12.5px; }
  .file .x { color: var(--muted); font-size: 22px; line-height: 1; margin-left: 6px; background: none; border: 0; cursor: pointer; }
  .go { background: var(--accent); color: #fff; border: 0; border-radius: 16px; font: inherit; font-weight: 700; font-size: 19px; cursor: pointer; padding: 22px; box-shadow: 0 8px 20px rgba(15,107,92,.28); }
  .go small { display: block; font-weight: 500; font-size: 13px; opacity: .85; margin-top: 6px; }
  .go:disabled { opacity: .45; cursor: not-allowed; box-shadow: none; }
  .privacy { color: var(--muted); font-size: 13px; margin-top: 12px; }

  .error { background: #fbe9e6; color: #8f2f20; border: 1px solid #f0c4bc; border-radius: 12px; padding: 14px 18px; margin-top: 18px; }
  .error[hidden], #results[hidden] { display: none; }

  .status { display: flex; align-items: center; flex-wrap: wrap; gap: 12px; margin: 26px 0 16px; }
  .badge { background: var(--accent-soft); color: var(--accent); font-weight: 600; padding: 6px 14px; border-radius: 999px; font-size: 13.5px; }
  .badge.warn { background: var(--warn-soft); color: var(--warn); }
  .status span.t { color: var(--muted); }
  .warnbox { background: var(--warn-soft); color: var(--warn); border-radius: 12px; padding: 12px 18px; margin-bottom: 16px; font-size: 14px; }

  .cards { display: grid; grid-template-columns: repeat(auto-fit, minmax(190px, 1fr)); gap: 16px; }
  .c { background: var(--card); border: 1px solid var(--line); border-radius: 14px; padding: 18px 20px; }
  .c label { color: var(--muted); font-size: 13px; font-weight: 500; }
  .c div { font-size: 24px; font-weight: 700; margin-top: 6px; letter-spacing: -0.5px; font-variant-numeric: tabular-nums; }
  .c.cr div { color: var(--credit); } .c.db div { color: var(--debit); }
  .c.big { background: var(--ink); border-color: var(--ink); }
  .c.big label { color: #aab2c5; } .c.big div { color: #fff; }

  .panel { background: var(--card); border: 1px solid var(--line); border-radius: 16px; margin-top: 26px; overflow: hidden; }
  .tabs { display: flex; padding: 14px 12px 0; border-bottom: 1px solid var(--line); background: #faf8f5; overflow-x: auto; }
  .tab { font: inherit; padding: 12px 9px; border-radius: 10px 10px 0 0; color: var(--muted); font-weight: 600; font-size: 14px; white-space: nowrap; background: none; border: 1px solid transparent; border-bottom: 0; cursor: pointer; }
  .tab em { font-style: normal; background: var(--line); color: var(--muted); font-size: 11.5px; padding: 2px 6px; border-radius: 999px; margin-left: 4px; }
  .tab.on { background: var(--card); color: var(--accent); box-shadow: inset 0 3px 0 var(--accent); border-color: var(--line); margin-bottom: -1px; padding-bottom: 13px; }
  .tab.on em { background: var(--accent-soft); color: var(--accent); }
  .tools { display: flex; justify-content: space-between; align-items: center; padding: 18px 24px; flex-wrap: wrap; gap: 12px; }
  .tools h4 { font-size: 17px; }
  .btns { display: flex; gap: 10px; }
  .btn { font: inherit; border: 1px solid var(--line); border-radius: 10px; padding: 9px 16px; font-weight: 600; font-size: 13.5px; background: #fff; color: var(--ink); text-decoration: none; cursor: pointer; }
  .btn.p { background: var(--accent-soft); color: var(--accent); border-color: #c5e0da; }
  .tablewrap { overflow-x: auto; }
  table { width: 100%; border-collapse: collapse; }
  th { text-align: left; font-size: 12.5px; text-transform: uppercase; letter-spacing: .6px; color: var(--muted); padding: 12px 24px; background: #faf8f5; border-top: 1px solid var(--line); border-bottom: 1px solid var(--line); font-weight: 600; white-space: nowrap; }
  td { padding: 13px 24px; border-bottom: 1px solid #f0ece5; font-variant-numeric: tabular-nums; }
  th.r, td.r { text-align: right; }
  td.d { color: var(--debit); font-weight: 600; } td.cr { color: var(--credit); font-weight: 600; } td.z { color: #b9bdc7; }
  tr.tot td { font-weight: 700; background: #faf8f5; border-bottom: 0; }
  td.empty { text-align: center; color: var(--muted); padding: 28px; }

  @media (max-width: 820px) { .upload { grid-template-columns: 1fr; } .wrap { padding: 0 16px 40px; } h1 { font-size: 26px; } }
</style>
</head>
<body>
  <div class="topbar"><div class="logo"><i></i>Balance Sheet Builder</div></div>

  <div class="wrap">
    <h1>Turn your bank statement into a balance sheet</h1>
    <div class="sub">Upload a PDF statement. Every transaction is checked against the running balance.</div>

    <div class="upload">
      <label class="drop" id="drop">
        <input type="file" id="file" accept="application/pdf,.pdf" hidden>
        <div class="ico"></div>
        <div>
          <h3>Drag &amp; drop your statement here</h3>
          <p>or click to browse &middot; PDF, up to 20 MB</p>
        </div>
        <div class="file" id="chip" hidden>
          <span class="pdf">PDF</span>
          <div><span id="fname"></span><small id="fsize"></small></div>
          <button class="x" id="clear" type="button" title="Remove file">&times;</button>
        </div>
      </label>
      <button class="go" id="go" type="button" disabled>Show Balance Sheet<small>Takes a few seconds</small></button>
    </div>
    <div class="privacy">Your statement is read in memory and is not saved. The Excel download is kept for 15 minutes, then discarded.</div>
    <div class="error" id="err" hidden></div>

    <div id="results" hidden>
      <div class="status"><span class="badge" id="badge"></span><span class="t" id="statusText"></span></div>
      <div class="warnbox" id="warnbox" hidden></div>
      <div class="cards" id="cards"></div>
      <div class="panel">
        <div class="tabs" id="tabs" role="tablist"></div>
        <div id="panelBody"></div>
      </div>
    </div>
  </div>

<script>
const $ = id => document.getElementById(id);
const fmt = n => n.toLocaleString('en-US', {minimumFractionDigits: 2, maximumFractionDigits: 2});
const MAX_BYTES = 20 * 1024 * 1024;
let chosen = null, state = null, tabs = [], active = 'monthly', monthFilter = '';

function el(tag, cls, text) {
  const e = document.createElement(tag);
  if (cls) e.className = cls;
  if (text !== undefined) e.textContent = text;
  return e;
}
function showError(msg) { const e = $('err'); e.textContent = msg; e.hidden = false; }
function clearError() { $('err').hidden = true; }
function niceDate(iso) {
  const m = ['Jan','Feb','Mar','Apr','May','Jun','Jul','Aug','Sep','Oct','Nov','Dec'];
  const p = iso.split('-');
  return p[2] + ' ' + m[parseInt(p[1], 10) - 1] + ' ' + p[0];
}

function setFile(f) {
  clearError();
  if (!f) return;
  if (!(f.type === 'application/pdf' || /\.pdf$/i.test(f.name))) { showError('Please choose a PDF file.'); return; }
  if (f.size > MAX_BYTES) { showError('That file is larger than 20 MB.'); return; }
  chosen = f;
  $('fname').textContent = f.name;
  $('fsize').textContent = (f.size / 1024).toFixed(0) + ' KB';
  $('chip').hidden = false;
  $('go').disabled = false;
}
function clearFile() {
  chosen = null; $('file').value = '';
  $('chip').hidden = true; $('go').disabled = true;
}

$('file').addEventListener('change', e => setFile(e.target.files[0]));
$('clear').addEventListener('click', e => { e.preventDefault(); e.stopPropagation(); clearFile(); });
const drop = $('drop');
['dragenter', 'dragover'].forEach(ev => drop.addEventListener(ev, e => { e.preventDefault(); drop.classList.add('over'); }));
['dragleave', 'drop'].forEach(ev => drop.addEventListener(ev, e => { e.preventDefault(); drop.classList.remove('over'); }));
drop.addEventListener('drop', e => setFile(e.dataTransfer.files[0]));

$('go').addEventListener('click', async () => {
  if (!chosen) return;
  clearError();
  const btn = $('go');
  btn.disabled = true;
  btn.firstChild.textContent = 'Reading statement...';
  try {
    const res = await fetch('/api/process', {method: 'POST', headers: {'Content-Type': 'application/pdf'}, body: chosen});
    if (!res.ok) {
      let msg = 'Something went wrong (' + res.status + ').';
      try { const j = await res.json(); if (j.detail) msg = j.detail; } catch (e) {}
      showError(msg);
    } else {
      state = await res.json();
      monthFilter = ''; active = 'monthly';
      render();
    }
  } catch (e) {
    showError('Could not reach the server: ' + e);
  } finally {
    btn.firstChild.textContent = 'Show Balance Sheet';
    btn.disabled = !chosen;
  }
});

function render() {
  const s = state.summary;
  $('results').hidden = false;

  const badge = $('badge');
  if (s.reconciled) { badge.textContent = '✓ Reconciled'; badge.className = 'badge'; }
  else { badge.textContent = '⚠ ' + s.unexplained.length + ' row(s) need checking'; badge.className = 'badge warn'; }
  $('statusText').textContent = s.count + ' transactions read from ' + s.pages + ' page' + (s.pages === 1 ? '' : 's') +
    ' · ' + niceDate(s.period_from) + ' – ' + niceDate(s.period_to);

  const wb = $('warnbox');
  if (s.unexplained.length) {
    wb.hidden = false;
    wb.textContent = 'These rows do not match the running balance, so check them against your PDF: ' +
      s.unexplained.slice(0, 8).map(u => 'S.No ' + u.sno + ' (off by ' + fmt(u.difference) + ')').join('; ') +
      (s.unexplained.length > 8 ? '; ...' : '');
  } else { wb.hidden = true; }

  const cards = $('cards'); cards.replaceChildren();
  [['Opening balance', s.opening, ''], ['Total credits', s.credits, 'cr'], ['Total debits', s.debits, 'db'],
   ['WHT deducted (20%)', s.wht, ''], ['Closing balance', s.closing, 'big']].forEach(([label, v, cls]) => {
    const c = el('div', 'c ' + cls);
    c.appendChild(el('label', '', label));
    c.appendChild(el('div', '', fmt(v)));
    cards.appendChild(c);
  });

  tabs = [{id: 'monthly', label: 'Monthly', count: state.monthly.length, kind: 'monthly'},
          {id: 'summary', label: 'Summary', count: state.categories.length, kind: 'summary'}];
  state.categories.forEach(c => tabs.push({id: 'cat:' + c.name, label: c.name, count: c.count, kind: 'tx', cat: c.name}));
  tabs.push({id: 'all', label: 'All', count: state.transactions.length, kind: 'tx', cat: null});

  renderTabs(); renderPanel();
  $('results').scrollIntoView({behavior: 'smooth', block: 'start'});
}

function renderTabs() {
  const box = $('tabs'); box.replaceChildren();
  tabs.forEach(t => {
    const b = el('button', 'tab' + (t.id === active ? ' on' : ''), t.label);
    b.type = 'button'; b.setAttribute('role', 'tab');
    b.appendChild(el('em', '', String(t.count)));
    b.addEventListener('click', () => { active = t.id; renderTabs(); renderPanel(); });
    box.appendChild(b);
  });
}

function buildTable(cols, rows, totals) {
  const wrap = el('div', 'tablewrap'), t = el('table'), thead = el('thead'), hr = el('tr');
  cols.forEach(c => hr.appendChild(el('th', c.right ? 'r' : '', c.label)));
  thead.appendChild(hr); t.appendChild(thead);
  const tb = el('tbody');
  if (!rows.length) {
    const tr = el('tr'), td = el('td', 'empty', 'No transactions'); td.colSpan = cols.length;
    tr.appendChild(td); tb.appendChild(tr);
  }
  rows.forEach(r => {
    const tr = el('tr');
    cols.forEach(c => {
      const v = r[c.key];
      let cls = c.right ? 'r' : '';
      if (c.kind === 'debit') cls += v > 0 ? ' d' : ' z';
      if (c.kind === 'credit') cls += v > 0 ? ' cr' : ' z';
      if (c.kind === 'wht') cls += v > 0 ? '' : ' z';
      tr.appendChild(el('td', cls.trim(), c.num ? fmt(v) : String(v)));
    });
    tb.appendChild(tr);
  });
  if (totals && rows.length) {
    const tr = el('tr', 'tot');
    cols.forEach(c => {
      const v = totals[c.key];
      tr.appendChild(el('td', c.right ? 'r' : '', v === undefined ? '' : (c.num ? fmt(v) : String(v))));
    });
    tb.appendChild(tr);
  }
  t.appendChild(tb); wrap.appendChild(t);
  return wrap;
}

const sum = (rows, k) => rows.reduce((a, r) => a + r[k], 0);

function renderPanel() {
  const body = $('panelBody'); body.replaceChildren();
  const tab = tabs.find(t => t.id === active);

  const tools = el('div', 'tools');
  tools.appendChild(el('h4', '', tab.label));
  const btns = el('div', 'btns');
  if (tab.kind === 'tx') {
    const sel = el('select', 'btn');
    const all = el('option', '', 'All months'); all.value = ''; sel.appendChild(all);
    state.monthly.forEach(m => { const o = el('option', '', m.month); o.value = m.month; sel.appendChild(o); });
    sel.value = monthFilter;
    sel.addEventListener('change', () => { monthFilter = sel.value; renderPanel(); });
    btns.appendChild(sel);
  }
  const dl = el('a', 'btn p', 'Download Excel');
  dl.href = '/api/download/' + encodeURIComponent(state.token);
  btns.appendChild(dl);
  tools.appendChild(btns);
  body.appendChild(tools);

  if (tab.kind === 'monthly') {
    const rows = state.monthly;
    body.appendChild(buildTable([
      {key: 'month', label: 'Month'},
      {key: 'gross_profit', label: 'Gross profit', right: true, num: true},
      {key: 'credits', label: 'Credits', right: true, num: true, kind: 'credit'},
      {key: 'debits', label: 'Debits', right: true, num: true, kind: 'debit'},
      {key: 'wht', label: 'WHT', right: true, num: true, kind: 'wht'},
      {key: 'net', label: 'Net movement', right: true, num: true},
      {key: 'closing_balance', label: 'Closing balance', right: true, num: true}],
      rows, {month: 'Total', gross_profit: sum(rows, 'gross_profit'), credits: sum(rows, 'credits'),
             debits: sum(rows, 'debits'), wht: sum(rows, 'wht'), net: sum(rows, 'net')}));
  } else if (tab.kind === 'summary') {
    const rows = state.categories;
    body.appendChild(buildTable([
      {key: 'name', label: 'Category'},
      {key: 'count', label: 'Transactions', right: true},
      {key: 'debits', label: 'Debits', right: true, num: true, kind: 'debit'},
      {key: 'credits', label: 'Credits', right: true, num: true, kind: 'credit'},
      {key: 'wht', label: 'WHT', right: true, num: true, kind: 'wht'}],
      rows, {name: 'Total', count: sum(rows, 'count'), debits: sum(rows, 'debits'),
             credits: sum(rows, 'credits'), wht: sum(rows, 'wht')}));
  } else {
    const rows = state.transactions.filter(r =>
      (tab.cat === null || r.category === tab.cat) && (!monthFilter || r.date.startsWith(monthFilter)));
    const cols = [{key: 'date', label: 'Date'}, {key: 'desc', label: 'Description'}];
    if (tab.cat === null) cols.push({key: 'category', label: 'Category'});
    cols.push({key: 'debit', label: 'Debit', right: true, num: true, kind: 'debit'},
              {key: 'credit', label: 'Credit', right: true, num: true, kind: 'credit'},
              {key: 'wht', label: 'WHT', right: true, num: true, kind: 'wht'},
              {key: 'balance', label: 'Balance', right: true, num: true});
    body.appendChild(buildTable(cols, rows, {desc: 'Total (' + rows.length + ' transactions)',
      debit: sum(rows, 'debit'), credit: sum(rows, 'credit'), wht: sum(rows, 'wht')}));
  }
}
</script>
</body>
</html>
"""
